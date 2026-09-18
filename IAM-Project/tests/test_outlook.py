import io
import json
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from html import escape
from pathlib import Path
from threading import Event
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import app
from outlook_collector import OutlookCollector, resolve_folder, validate_options
from outlook_store import OutlookStore
from outlook_tables import extract_rows, matches_subject


def table(headers, *rows):
    return "<table>" + "".join(
        "<tr>" + "".join(f"<td><b>{escape(str(cell))}</b></td>" for cell in row) + "</tr>"
        for row in (headers, *rows)
    ) + "</table>"


HEADERS = ["Commande GPON", "ONT", "Version", "Technologie", "Intitulé client",
           "Login", "MSAN / SRO", "PCO", "brin", "catégorie", "Pose PCO(O/N)",
           "Pose Nouveau Splitter O/N", "GPS PCO", "GPS NV Splitter", "1FO"]
ROW = ["101000001 -- RECGPON", "ONT0001", "Autre", "HUAWEI", "Client exemple",
       "client.test", "GHI-FF-EXEMPLE", "5121/1", "2", "Façade", "N", "N",
       "34.0315747,-5.0644695", "N", "73M"]


def mailbox_tree():
    def folder(name, children=None):
        return SimpleNamespace(
            StoreID="store1", EntryID=name, FolderPath=name,
            Folders=SimpleNamespace(Item=Mock(side_effect=(children or {}).__getitem__)),
        )
    ftth = folder("Archivage/FTTH")
    archive = folder("Archivage", {"FTTH": ftth})
    inbox = folder("Réception")
    root = folder("Compte", {"Réception": inbox, "Archivage": archive})
    inbox.Parent = root
    namespace = SimpleNamespace(
        GetDefaultFolder=Mock(return_value=inbox),
        Folders=SimpleNamespace(Item=Mock(side_effect={"Compte": root}.__getitem__)),
    )
    return namespace, inbox, archive, ftth


class TableTests(unittest.TestCase):
    def test_subjects_ignore_accents_case_and_extra_spaces(self):
        for subject in ["RE: ACTIVATION  commande FTTH client", "Création CD FTTH",
                        "TR: CREATION GPON", "commandes GPON", "Activation Mise en service"]:
            self.assertTrue(matches_subject(subject))
        self.assertFalse(matches_subject("Compte rendu de réunion"))

    def test_first_example_multiple_rows_and_nested_outlook_layout(self):
        second = [*ROW]
        second[0] = "101000002 -- RECGPON"
        second[12] = "34,031897 /-5,039591"
        html = "<table><tr><td>Bonjour" + table(HEADERS, ROW, second) + "</td></tr></table>"
        rows = extract_rows(html)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["odf"], "GHI-FF-EXEMPLE")
        self.assertEqual(rows[0]["msan"], "")
        self.assertEqual(rows[0]["longueur"], "73M")
        self.assertEqual(rows[1]["gps_pco"], "34,031897 /-5,039591")

    def test_second_example_separates_msan_odf_and_combines_coordinates(self):
        headers = ["COM", "ONT", "NOM DU CLIENT", "MSAN", "ODF", "NOUVEAU PCO",
                   "LONGEUR", "PCO", "TYPE DE PCO", "num de Brin", "CGPS PCO x",
                   "CGPS PCO y", "login", "Pose nouveau splitter O/N"]
        row = ["101000003", "ONT0003", "Client test", "GHI\nAIN   CHEGAGE", "OFCH",
               "OUI", "61", "911/2", "FACADE", "5", "33.889088", "-5.038047", "LOGIN3", "N"]
        actual = extract_rows(table(headers, row))[0]
        self.assertEqual(actual["odf"], "OFCH")
        self.assertEqual(actual["msan"], "GHI AIN CHEGAGE")
        self.assertEqual(actual["gps_pco"], "33.889088 / -5.038047")
        self.assertEqual(actual["brin"], "5")
        self.assertEqual(actual["pose_pco"], "OUI")

    def test_msan_olt_alias_stays_separate_from_odf(self):
        actual = extract_rows(table(
            ["COM", "Login", "MSAN / SRO", "MSAN OLT"],
            ["CMD1", "LOGIN1", "ODF1", "OLT1"],
        ))[0]
        self.assertEqual(actual["odf"], "ODF1")
        self.assertEqual(actual["msan"], "OLT1")

    def test_third_example_keeps_dfoi_command_and_missing_columns_blank(self):
        row = [*ROW[:13]]
        row[0], row[4], row[7], row[12] = "DFOIWC00000001", "Exemple", "2. 712/2", "34.026322\n-4.987193"
        headers = [*HEADERS[:13]]
        headers[4] = "Nom client"
        actual = extract_rows(table(headers, row))[0]
        self.assertEqual(actual["commande"], "DFOIWC00000001")
        self.assertEqual(actual["gps_pco"], "34.026322 -4.987193")
        self.assertEqual(actual["pco"], "2. 712/2")
        self.assertEqual(actual["longueur"], "")

    def test_merged_cells_repeated_headers_and_unrelated_tables(self):
        html = '''<table><tr><th>Commande GPON</th><th>Login</th><th>ODF</th></tr>
          <tr><td>CMD1</td><td>LOGIN1</td><td rowspan="2">ODF1</td></tr>
          <tr><td>CMD2</td><td>LOGIN2</td></tr></table>'''
        html += table(["Signature", "Téléphone", "Adresse"], ["Nom", "0000", "Ville"])
        html += table(HEADERS, ROW, HEADERS, ROW)  # repeated quote is not added twice
        rows = extract_rows(html)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1]["odf"], "ODF1")


class StoreTests(unittest.TestCase):
    def test_legacy_rows_are_reparsed_without_duplicates_or_data_loss(self):
        with tempfile.TemporaryDirectory() as directory:
            store = OutlookStore(Path(directory))
            old_row = {"commande": "CMD1", "login": "LOGIN1", "odf": "OLT1 | ODF1"}
            store.record("old-mail", {}, [old_row])
            db = store.connect()
            try:
                with db:
                    db.execute("UPDATE messages SET metadata=? WHERE id=?", (json.dumps({}), "old-mail"))
            finally:
                db.close()
            self.assertFalse(store.seen("old-mail"))
            with self.assertRaises(ValueError):
                store.record("old-mail", {}, [])
            self.assertEqual(store.summary()["rows"][0]["odf"], "OLT1 | ODF1")
            fresh_rows = extract_rows(table(
                ["COM", "Login", "MSAN OLT", "ODF"], ["CMD1", "LOGIN1", "OLT1", "ODF1"],
            ))
            self.assertTrue(store.record("old-mail", {}, fresh_rows))
            self.assertTrue(store.seen("old-mail"))
            self.assertFalse(store.record("old-mail", {}, fresh_rows))
            summary = store.summary()
            self.assertEqual((summary["emails"], summary["total"]), (1, 1))
            self.assertEqual(summary["rows"][0]["odf"], "ODF1")
            self.assertEqual(summary["rows"][0]["msan"], "OLT1")
            book = load_workbook(io.BytesIO(store.excel_bytes()))
            headings = {cell.value: cell.column for cell in book.active[1]}
            self.assertEqual(book.active.cell(2, headings["ODF"]).value, "ODF1")
            self.assertEqual(book.active.cell(2, headings["MSAN"]).value, "OLT1")
            book.close()

    def test_dedup_persists_across_restarts_and_excel_cells_are_text(self):
        with tempfile.TemporaryDirectory() as directory:
            store = OutlookStore(Path(directory))
            rows = extract_rows(table(HEADERS, ROW))
            rows[0]["client"] = '=HYPERLINK("https://example.invalid")'
            rows[0]["ont"] = "000123"
            self.assertTrue(store.record("mail1", {"subject": "Test", "sender": "test@example.invalid"}, rows))
            store = OutlookStore(Path(directory))
            self.assertTrue(store.seen("mail1"))
            self.assertFalse(store.record("mail1", {}, rows))
            self.assertEqual(store.summary()["total"], 1)
            store.write_excel()
            book = load_workbook(io.BytesIO(store.excel_path.read_bytes()))
            sheet = book.active
            headings = {cell.value: cell.column for cell in sheet[1]}
            self.assertIn("ODF", headings)
            self.assertEqual(headings["MSAN"], headings["Longueur"] + 1)
            self.assertEqual(sheet.cell(2, headings["Client"]).data_type, "s")
            self.assertEqual(sheet.cell(2, headings["ONT"]).value, "000123")
            self.assertEqual(sheet.max_row, 2)
            book.close()

    def test_unrecognized_email_is_counted_without_empty_excel_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            store = OutlookStore(Path(directory))
            store.record("mail-empty", {"subject": "CREATION GPON"}, [])
            summary = store.summary()
            self.assertEqual((summary["emails"], summary["skipped"], summary["total"]), (1, 1, 0))


class CollectionTests(unittest.TestCase):
    def test_worker_starts_once_stops_and_releases_com_on_its_thread(self):
        pythoncom = ModuleType("pythoncom")
        pythoncom.CoInitialize = Mock()
        pythoncom.CoUninitialize = Mock()
        win32com = ModuleType("win32com")
        client_module = ModuleType("win32com.client")
        client_module.Dispatch = Mock()
        win32com.client = client_module
        folder = client_module.Dispatch.return_value.GetNamespace.return_value.GetDefaultFolder.return_value
        folder.FolderPath = "Inbox"
        entered, release = Event(), Event()

        def scan(_folder, _cutoff):
            entered.set()
            release.wait(3)
            return 0

        with tempfile.TemporaryDirectory() as directory:
            collector = OutlookCollector(OutlookStore(Path(directory)))
            modules = {"pythoncom": pythoncom, "win32com": win32com, "win32com.client": client_module}
            with patch.dict("sys.modules", modules), patch("outlook_collector.sys.platform", "win32"), \
                    patch.object(collector, "scan", side_effect=scan):
                try:
                    collector.start({})
                    self.assertTrue(entered.wait(2))
                    with self.assertRaisesRegex(ValueError, "déjà active"):
                        collector.start({})
                    collector.stop()
                    self.assertEqual(collector.status()["status"], "STOPPING")
                finally:
                    collector.stop()
                    release.set()
                    if collector.thread:
                        collector.thread.join(3)
                self.assertFalse(collector.thread.is_alive())
                self.assertEqual(collector.status()["status"], "STOPPED")
                pythoncom.CoInitialize.assert_called_once()
                pythoncom.CoUninitialize.assert_called_once()
                self.assertTrue(collector.store.excel_path.exists())

    def test_inbox_and_user_subfolder_resolution(self):
        namespace = Mock()
        inbox = namespace.GetDefaultFolder.return_value
        self.assertIs(resolve_folder(namespace, ""), inbox)
        nested = resolve_folder(namespace, "FTTH/Activations")
        self.assertIs(nested, inbox.Folders.Item.return_value.Folders.Item.return_value)
        namespace.GetDefaultFolder.assert_called_with(6)
        inbox.Folders.Item.assert_called_with("FTTH")
        inbox.Folders.Item.return_value.Folders.Item.assert_called_with("Activations")

    def test_archive_beside_inbox_and_full_mailbox_paths(self):
        namespace, inbox, archive, ftth = mailbox_tree()
        self.assertIs(resolve_folder(namespace, "Archivage"), archive)
        self.assertIs(resolve_folder(namespace, "Archivage/FTTH"), ftth)
        self.assertIs(resolve_folder(namespace, "Compte/Archivage"), archive)
        self.assertIs(resolve_folder(namespace, "\\\\Compte\\Archivage"), archive)
        self.assertIs(resolve_folder(namespace, "Compte/Réception"), inbox)
        with self.assertRaisesRegex(ValueError, "Dossier Outlook introuvable"):
            resolve_folder(namespace, "Absent")

    def test_scan_watches_inbox_and_archive_without_scanning_subfolders(self):
        namespace, inbox, archive, _ftth = mailbox_tree()
        cutoff = datetime(2026, 1, 2, tzinfo=timezone.utc)
        with tempfile.TemporaryDirectory() as directory:
            collector = OutlookCollector(OutlookStore(Path(directory)))
            with patch.object(collector, "scan", side_effect=[2, 3]) as scan:
                self.assertEqual(collector.scan_folders(namespace, "Archivage", cutoff), 5)
                self.assertEqual([call.args[0] for call in scan.call_args_list], [inbox, archive])
            self.assertEqual(collector.status()["folder_label"], "Réception + Archivage")
            self.assertIsNone(collector.status()["error"])

    def test_invalid_extra_folder_does_not_stop_inbox(self):
        namespace, inbox, _archive, _ftth = mailbox_tree()
        cutoff = datetime(2026, 1, 2, tzinfo=timezone.utc)
        with tempfile.TemporaryDirectory() as directory:
            collector = OutlookCollector(OutlookStore(Path(directory)))
            with patch.object(collector, "scan", return_value=1) as scan:
                self.assertEqual(collector.scan_folders(namespace, "Absent", cutoff), 1)
                scan.assert_called_once_with(inbox, cutoff)
            self.assertIn("Dossier Outlook introuvable", collector.status()["error"])

    def test_folder_failure_does_not_prevent_scanning_the_other_folder(self):
        namespace, inbox, archive, _ftth = mailbox_tree()
        cutoff = datetime(2026, 1, 2, tzinfo=timezone.utc)
        with tempfile.TemporaryDirectory() as directory:
            collector = OutlookCollector(OutlookStore(Path(directory)))
            with patch.object(collector, "scan", side_effect=[RuntimeError("offline"), 1]) as scan:
                self.assertEqual(collector.scan_folders(namespace, "Archivage", cutoff), 1)
                self.assertEqual([call.args[0] for call in scan.call_args_list], [inbox, archive])
            self.assertIn("Lecture impossible", collector.status()["error"])

    def test_explicit_inbox_path_is_scanned_only_once(self):
        namespace, inbox, _archive, _ftth = mailbox_tree()
        cutoff = datetime(2026, 1, 2, tzinfo=timezone.utc)
        with tempfile.TemporaryDirectory() as directory:
            collector = OutlookCollector(OutlookStore(Path(directory)))
            with patch.object(collector, "scan", return_value=1) as scan:
                self.assertEqual(collector.scan_folders(namespace, "Compte/Réception", cutoff), 1)
                scan.assert_called_once_with(inbox, cutoff)

    def test_scan_filters_dates_subjects_and_deduplicates(self):
        cutoff = datetime(2026, 1, 2, tzinfo=timezone.utc)
        messages = []
        for i, (subject, received) in enumerate([
            ("CREATION GPON", cutoff + timedelta(hours=1)),
            ("Autre sujet", cutoff + timedelta(hours=1)),
            ("CREATION GPON", cutoff - timedelta(seconds=1)),
        ]):
            accessor = Mock()
            accessor.GetProperty.return_value = f"message-{i}"
            messages.append(SimpleNamespace(Class=43, Subject=subject, ReceivedTime=received,
                PropertyAccessor=accessor, EntryID=f"entry-{i}", HTMLBody=table(HEADERS, ROW),
                SenderEmailAddress="test@example.invalid", SenderName="Test"))
        items = Mock()
        items.Count = len(messages)
        items.Item.side_effect = lambda index: messages[index - 1]
        folder = Mock(StoreID="store", FolderPath="Inbox")
        folder.Items.Restrict.return_value = items
        with tempfile.TemporaryDirectory() as directory:
            collector = OutlookCollector(OutlookStore(Path(directory)))
            self.assertEqual(collector.scan(folder, cutoff), 1)
            self.assertEqual(collector.scan(folder, cutoff), 0)
            self.assertEqual(collector.status()["total"], 1)
            collector.stop_event.set()
            self.assertEqual(collector.scan(folder, cutoff), 0)

    def test_option_validation(self):
        self.assertEqual(validate_options({})["folder"], "")
        self.assertEqual(validate_options({"folder": "FTTH\\Activations"})["folder"], "FTTH/Activations")
        for payload in [{"mode": "invalid"}, {"mode": "since", "since": "bad"},
                        {"interval": 0}, {"folder": "FTTH/../Other"}]:
            with self.assertRaises(ValueError):
                validate_options(payload)

    def test_routes_use_independent_collector_and_protected_writes(self):
        with tempfile.TemporaryDirectory() as directory:
            store = OutlookStore(Path(directory))
            collector = OutlookCollector(store)
            with patch.object(app, "OUTLOOK_STORE", store), patch.object(app, "OUTLOOK_COLLECTOR", collector):
                client = app.app.test_client()
                self.assertEqual(client.get("/api/outlook").status_code, 200)
                self.assertEqual(client.post("/api/outlook/start", json={}).status_code, 403)
                response = client.get("/api/outlook/result.xlsx")
                self.assertEqual(response.status_code, 200)
                book = load_workbook(io.BytesIO(response.data))
                self.assertEqual(book.active.max_row, 1)
                book.close()
                with patch.object(collector, "start") as start:
                    with patch.dict(app.jobs, {"busy": {"kind": "RENSEIGNER", "status": "RUNNING"}}):
                        response = client.post("/api/outlook/start", json={"folder": "FTTH"}, headers={"X-Requested-With": "FB-EMM"})
                    self.assertEqual(response.status_code, 200)
                    start.assert_called_once_with({"folder": "FTTH"})


if __name__ == "__main__":
    unittest.main()
