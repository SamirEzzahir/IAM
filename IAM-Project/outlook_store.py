"""Durable mail deduplication and Excel snapshots for Outlook collection."""

from __future__ import annotations

import io
import json
import sqlite3
from contextlib import closing
from pathlib import Path
from threading import Lock

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from outlook_tables import COLUMNS, PARSER_VERSION


META_COLUMNS = (
    ("received", "Date réception"), ("sender", "Expéditeur"),
    ("subject", "Sujet"), ("folder", "Dossier Outlook"),
)


class OutlookStore:
    def __init__(self, directory: Path):
        self.directory = directory
        self.path = directory / "outlook.sqlite3"
        self.excel_path = directory / "collecte_outlook.xlsx"
        self.lock = Lock()

    def connect(self):
        self.directory.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(self.path, timeout=15)
        db.executescript("""
            CREATE TABLE IF NOT EXISTS messages (
                id TEXT PRIMARY KEY, metadata TEXT NOT NULL, row_count INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS extracted_rows (
                id INTEGER PRIMARY KEY, message_id TEXT NOT NULL, payload TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS settings (id INTEGER PRIMARY KEY, payload TEXT NOT NULL);
        """)
        return db

    def seen(self, message_id: str) -> bool:
        with self.lock, closing(self.connect()) as db:
            row = db.execute("SELECT metadata FROM messages WHERE id=?", (message_id,)).fetchone()
        return bool(row and json.loads(row[0]).get("parser_version", 1) >= PARSER_VERSION)

    def record(self, message_id: str, metadata: dict, rows: list[dict]) -> bool:
        with self.lock, closing(self.connect()) as db, db:
            previous = db.execute("SELECT metadata, row_count FROM messages WHERE id=?", (message_id,)).fetchone()
            if previous:
                if json.loads(previous[0]).get("parser_version", 1) >= PARSER_VERSION:
                    return False
                if previous[1] and not rows:
                    # Keep historical data if the original table cannot be read.
                    raise ValueError("Le tableau d’origine n’est plus lisible ; anciennes lignes conservées.")
                # Reparse an old message atomically instead of guessing how its
                # previously merged ODF/MSAN values should be split.
                db.execute("DELETE FROM extracted_rows WHERE message_id=?", (message_id,))
            metadata = {**metadata, "parser_version": PARSER_VERSION}
            db.execute("INSERT OR REPLACE INTO messages VALUES (?, ?, ?)", (
                message_id, json.dumps(metadata, ensure_ascii=False), len(rows),
            ))
            db.executemany("INSERT INTO extracted_rows(message_id, payload) VALUES (?, ?)", [
                (message_id, json.dumps({**row, **metadata}, ensure_ascii=False)) for row in rows
            ])
        return True

    def settings(self) -> dict:
        with self.lock, closing(self.connect()) as db:
            row = db.execute("SELECT payload FROM settings WHERE id=1").fetchone()
        return json.loads(row[0]) if row else {}

    def save_settings(self, value: dict) -> None:
        with self.lock, closing(self.connect()) as db, db:
            db.execute("INSERT OR REPLACE INTO settings VALUES (1, ?)", (json.dumps(value),))

    def summary(self) -> dict:
        with self.lock, closing(self.connect()) as db:
            emails, skipped, total = db.execute(
                "SELECT COUNT(*), COALESCE(SUM(row_count=0),0), COALESCE(SUM(row_count),0) FROM messages"
            ).fetchone()
            latest = db.execute("SELECT payload FROM extracted_rows ORDER BY id DESC LIMIT 100").fetchall()
        return {"emails": emails, "skipped": skipped, "total": total,
                "rows": [json.loads(item[0]) for item in latest]}

    def excel_bytes(self) -> bytes:
        with self.lock, closing(self.connect()) as db:
            rows = db.execute("SELECT payload FROM extracted_rows ORDER BY id").fetchall()
        book = Workbook()
        sheet = book.active
        sheet.title = "Collecte Outlook"
        columns = (*COLUMNS, *META_COLUMNS)
        sheet.append([label for _key, label in columns])
        for row_index, item in enumerate(rows, 2):
            row = json.loads(item[0])
            # Mail values are text, never Excel formulas; also retain leading zeros.
            sheet.append([ILLEGAL_CHARACTERS_RE.sub("", str(row.get(key, "")))[:32767] for key, _ in columns])
            for column in range(1, len(columns) + 1):
                sheet.cell(row_index, column).data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="B51E32")
            sheet.column_dimensions[cell.column_letter].width = 23
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    def write_excel(self) -> None:
        content = self.excel_bytes()
        temporary = self.excel_path.with_suffix(".tmp")
        temporary.write_bytes(content)
        temporary.replace(self.excel_path)
