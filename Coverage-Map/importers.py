"""Bounded, server-side imports. Never retain arbitrary source properties."""
import csv
import io
import json
import math
import re
import unicodedata
import zipfile

from defusedxml import ElementTree as ET
from openpyxl import load_workbook
from shapely.geometry import mapping, shape

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100000
CLIENT_COLUMNS = ("ODF", "Login", "Série ONT", "Nom Client", "Adresse Client", "NE", "PCO", "OLT")


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in text.lower() if not unicodedata.combining(c)))


def unzip_files(raw, extension):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        files = archive.infolist()
        if len(files) > 2000 or sum(f.file_size for f in files) > 100 * 1024 * 1024:
            raise ValueError("Archive trop volumineuse après décompression.")
        selected = [f for f in files if f.filename.lower().endswith(extension)]
        if not selected:
            raise ValueError("Aucun fichier compatible dans l'archive.")
        return [archive.read(f) for f in selected]


def coordinates(element):
    if element is None or not element.text:
        raise ValueError("Coordonnées KML manquantes.")
    return [[float(n) for n in token.split(",")[:2]] for token in element.text.split()]


def kml_features(raw):
    root = ET.fromstring(raw)
    # Strip namespaces after safe parsing; supports both namespaced and simple KML.
    for el in root.iter():
        el.tag = el.tag.split("}")[-1]
    result = []
    for pm in root.iter("Placemark"):
        name = pm.findtext("name") or "Sans nom"
        polygons = []
        for polygon in pm.iter("Polygon"):
            outer = coordinates(polygon.find("./outerBoundaryIs/LinearRing/coordinates"))
            holes = [coordinates(el) for el in polygon.findall("./innerBoundaryIs/LinearRing/coordinates")]
            polygons.append([outer, *holes])
        if polygons:
            result.append({"properties": {"name": name}, "geometry": {
                "type": "MultiPolygon", "coordinates": polygons}})
        for point in pm.iter("Point"):
            result.append({"properties": {"name": name}, "geometry": {
                "type": "Point", "coordinates": coordinates(point.find("coordinates"))[0]}})
    return result


def read_geo(raw, filename, points=False):
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "kmz":
        features = [f for document in unzip_files(raw, ".kml") for f in kml_features(document)]
    elif suffix == "kml":
        features = kml_features(raw)
    elif suffix in {"json", "geojson"}:
        document = json.loads(raw)
        if not isinstance(document, dict):
            raise ValueError("GeoJSON invalide.")
        features = document.get("features", []) if document.get("type") == "FeatureCollection" else [document]
    else:
        raise ValueError("Format attendu : KML, KMZ ou GeoJSON (WGS84).")
    result = []
    for feature in features:
        geometry = feature.get("geometry") or {}
        allowed = {"Point"} if points else {"Polygon", "MultiPolygon"}
        if geometry.get("type") not in allowed:
            continue
        obj = shape(geometry)
        if obj.is_empty or not obj.is_valid:
            raise ValueError("Géométrie vide ou invalide : corrigez le fichier avant l'import.")
        west, south, east, north = obj.bounds
        if not all(math.isfinite(v) for v in obj.bounds) or not (-180 <= west <= east <= 180 and -90 <= south <= north <= 90):
            raise ValueError("Les coordonnées doivent être en WGS84 (longitude, latitude).")
        props = feature.get("properties") or {}
        name = str(props.get("name") or props.get("Name") or props.get("PCO") or props.get("nom") or "Sans nom")[:250]
        result.append({"type": "Feature", "geometry": mapping(obj), "properties": {"name": name}})
        if len(result) > MAX_ROWS:
            raise ValueError("Trop d'objets dans ce fichier.")
    if not result:
        raise ValueError("Aucun point PCO trouvé." if points else "Aucun polygone trouvé.")
    return result


def read_table(raw, filename, kind):
    suffix = filename.lower().rsplit(".", 1)[-1]
    workbook = None
    if suffix == "xlsx":
        # Inspect uncompressed size before openpyxl processes the archive.
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Classeur trop volumineux après décompression.")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
    elif suffix == "csv":
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = raw.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = iter(csv.reader(io.StringIO(decoded), dialect))
    else:
        raise ValueError("Format attendu : .xlsx ou .csv (convertissez les anciens .xls).")
    try:
        headers = [normalize(v) for v in next(rows, [])]
        if not headers or len(headers) > 200:
            raise ValueError("En-têtes absents ou trop de colonnes.")
        required = {"pco", "nbreoccupe", "nbreport"} if kind == "occupation" else {"pco"}
        if not required.issubset(headers):
            raise ValueError("Colonnes requises : " + ("PCO, NBRE_OCCUPE, NBRE_PORT" if kind == "occupation" else "PCO"))
        result = []
        seen = set()
        for number, row in enumerate(rows, 2):
            if number > MAX_ROWS + 1:
                raise ValueError("Maximum 100 000 lignes par import.")
            if not any(value is not None and str(value).strip() for value in row):
                continue
            values = dict(zip(headers, row))
            if kind == "occupation":
                pco = str(values.get("pco") or "").strip()
                try:
                    used, total = float(values["nbreoccupe"]), float(values["nbreport"])
                    if not (0 <= used <= total) or not used.is_integer() or not total.is_integer():
                        raise ValueError()
                except (ValueError, TypeError):
                    raise ValueError(f"Occupation invalide à la ligne {number}.") from None
                if not normalize(pco) or normalize(pco) in seen:
                    raise ValueError(f"PCO vide ou en double à la ligne {number}.")
                seen.add(normalize(pco))
                result.append({"pco": pco, "used": int(used), "total": int(total), "free": int(total - used)})
            else:
                result.append({key: str(values.get(normalize(key)) or "").strip()[:2000] for key in CLIENT_COLUMNS})
        if not result:
            raise ValueError("Le fichier ne contient aucune ligne exploitable.")
        return result
    finally:
        if workbook:
            workbook.close()
